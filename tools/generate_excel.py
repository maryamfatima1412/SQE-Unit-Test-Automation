"""Generate docs/Unit_Test_Cases.xlsx  (Task 3: test cases in an Excel sheet + external CSV data).

Run from the repository root:   python tools/generate_excel.py

Sheets
  Read Me             what the workbook is and how it is used
  Data Driven Cases   >= 4 test cases; every case points to a CSV file in tests/data.
                      tests/data_driven/test_excel_driven.py READS this sheet and runs every CSV row.
  Test Inventory      every automated test function (built from the source code, so it is always true)
  Coverage Check      formulas that check "5 to 10 tests per person for each function / class / module"
  Known Defects       the xfail tests with their defect IDs
"""
import ast
import csv
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Unit_Test_Cases.xlsx"
DATA = ROOT / "tests" / "data"

MARYAM, HUDA = "Maryam Fatima", "Huda Asif"

# ---------------------------------------------------------------------------------------------
# Data-driven test cases (consumed by tests/data_driven/test_excel_driven.py)
# ---------------------------------------------------------------------------------------------
DATA_DRIVEN = [
    # id, module, function under test, handler, technique, description, csv file, columns, expected rule, owner, priority
    ("DD-01", "Module 1 - Grading", "grade_from_percentage", "grade_from_percentage", "BVA + ECP",
     "Every row of the CSV is a percentage; the grade returned must equal expected_grade (values on and next to each cut-off).",
     "grade_boundaries.csv", "percentage, expected_grade", "grade_from_percentage(percentage) == expected_grade", MARYAM, "High"),
    ("DD-02", "Module 1 - Grading", "grade_color", "grade_color", "ECP",
     "Every grade (and the '-' placeholder) must map to the expected colour constant (GREEN, BLUE, ORANGE or RED).",
     "grade_colors.csv", "grade, expected_color", "grade_color(grade) == student_portal.<expected_color>", HUDA, "Medium"),
    ("DD-03", "Module 3 - StudentDashboard", "StudentDashboard.overall_attendance", "overall_attendance", "ECP + BVA",
     "A pipe-separated list of statuses is turned into attendance records; total, present and percentage must match.",
     "attendance_cases.csv", "case_id, statuses, expected_total, expected_present, expected_percent",
     "overall_attendance(records) == (expected_total, expected_present, expected_percent)", HUDA, "High"),
    ("DD-04", "Module 3 - StudentDashboard", "StudentDashboard.average_grade", "average_grade", "ECP + BVA",
     "Result percentages and approval flags are turned into result records; grade, count and average must match.",
     "average_grade_cases.csv", "case_id, percentages, approved, expected_grade, expected_count, expected_average",
     "average_grade(records) == (expected_grade, expected_count, expected_average)", MARYAM, "High"),
    ("DD-05", "Module 2 - Storage", "Storage.set / Storage.get", "storage_round_trip", "ECP",
     "A JSON value is saved with set(), the application is 'restarted' (new Storage) and get() must return the same value.",
     "storage_roundtrip.csv", "case_id, key, json_value", "Storage().get(key) == json.loads(json_value) after set()", MARYAM, "Medium"),
]

# ---------------------------------------------------------------------------------------------
# Mapping from test class -> function / class under test
# ---------------------------------------------------------------------------------------------
MODULES = {"grading": "Module 1 - Grading", "storage": "Module 2 - Storage",
           "student_analytics": "Module 3 - StudentDashboard"}
UNDER_TEST = {
    "TestGradeFromPercentage": "grade_from_percentage()", "TestGradeColor": "grade_color()",
    "TestGradingModule": "MODULE: grading scale",
    "TestGet": "Storage.get()", "TestSet": "Storage.set()", "TestLogActivity": "Storage.log_activity()",
    "TestStorageClass": "CLASS: Storage",
    "TestOverallAttendance": "StudentDashboard.overall_attendance()", "TestAverageGrade": "StudentDashboard.average_grade()",
    "TestLoadCourseWiseAttendance": "StudentDashboard.load_course_wise_attendance()",
    "TestStudentDashboardClass": "CLASS: StudentDashboard",
}

F = "Arial"
HEAD = PatternFill("solid", fgColor="0A5F5C")
SUB = PatternFill("solid", fgColor="D5E8E7")
thin = Side(style="thin", color="9AA5A4")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="top", wrap_text=True)


def header(ws, row, cols):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row, i, h)
        c.font, c.fill, c.border = Font(name=F, bold=True, color="FFFFFF", size=10), HEAD, BRD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def put(ws, row, values, center=()):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font, c.border = Font(name=F, size=10), BRD
        c.alignment = CEN if i in center else WRAP


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def csv_rows(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return sum(1 for _ in csv.DictReader(f))


def collected_case_counts():
    """Number of executed cases (parametrised variants) per test function, from pytest itself."""
    out = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q", "-p", "no:cacheprovider", "tests/unit"],
                         cwd=ROOT, capture_output=True, text=True).stdout
    counts = Counter()
    for line in out.splitlines():
        if "::" in line:
            counts[re.sub(r"\[.*\]$", "", line.strip())] += 1
    return counts


def decorator_markers(fn):
    names = []
    for d in fn.decorator_list:
        node = d.func if isinstance(d, ast.Call) else d
        if isinstance(node, ast.Attribute) and isinstance(node.value, ast.Attribute) and node.value.attr == "mark":
            if node.attr == "parametrize":
                continue
            arg = ""
            if isinstance(d, ast.Call) and d.args and isinstance(d.args[0], ast.Constant):
                arg = f"={d.args[0].value}"
            names.append(node.attr + arg)
    return names


def build_inventory():
    counts = collected_case_counts()
    rows = []
    for path in sorted((ROOT / "tests" / "unit").rglob("test_*.py")):
        module = MODULES[path.parent.name]
        owner = MARYAM if path.stem.endswith("maryam") else HUDA
        tree = ast.parse(path.read_text(encoding="utf-8"))
        rel = path.relative_to(ROOT).as_posix()
        for cls in [n for n in tree.body if isinstance(n, ast.ClassDef) and n.name.startswith("Test")]:
            for fn in [n for n in cls.body if isinstance(n, ast.FunctionDef) and n.name.startswith("test_")]:
                doc = (ast.get_docstring(fn) or "").splitlines()[0]
                m = re.match(r"\[([^\]]+)\]\s*(.*)", doc)
                technique, text = (m.group(1), m.group(2)) if m else ("", doc)
                nodeid = f"{rel}::{cls.name}::{fn.name}"
                rows.append([None, module, UNDER_TEST[cls.name], fn.name, text, technique, owner,
                             counts.get(nodeid, 1), ", ".join(decorator_markers(fn)), rel])
    for i, r in enumerate(rows, 1):
        r[0] = f"UT-{i:03d}"
    return rows


def known_defects():
    rows = []
    tree = ast.parse((ROOT / "tests" / "known_defects" / "test_known_defects.py").read_text(encoding="utf-8"))
    for fn in [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name.startswith("test_")]:
        for d in fn.decorator_list:
            if isinstance(d, ast.Call) and getattr(d.func, "attr", "") == "xfail":
                reason = next(k.value.value for k in d.keywords if k.arg == "reason")
                rows.append([reason.split(":")[0], reason.split(":", 1)[1].strip(), fn.name,
                             "Assignment 2 defect confirmed at unit level" if reason.startswith("DEF") else "New defect found by unit tests"])
    return rows


def main():
    wb = Workbook()
    # ------------------------------------------------------------------ Read Me
    ws = wb.active
    ws.title = "Read Me"
    lines = [
        ("Unit Test Cases - Student Portal (Assignment 3: Build Unit Test Automation Solution)", True),
        ("Authors: Maryam Fatima (24F-3044) and Huda Asif (24F-3111)  |  Application under test: src/student_portal.py", False),
        ("", False),
        ("Sheet 'Data Driven Cases': the test cases that are executed from external CSV files (Task 3). The automated test "
         "tests/data_driven/test_excel_driven.py reads this sheet, opens the CSV named in column 'Test Data File' and runs every row.", False),
        ("Sheet 'Test Inventory': every automated unit test function with its module, function under test, technique and owner.", False),
        ("Sheet 'Coverage Check': formulas that verify the rule '5 to 10 tests per person for each function and for each class / module'.", False),
        ("Sheet 'Known Defects': tests that are marked xfail because they document a real defect.", False),
        ("", False),
        ("To add a data-driven test: add a row to 'Data Driven Cases' (and a CSV in tests/data), then add a handler with the "
         "same name in tests/data_driven/test_excel_driven.py. Re-generate this workbook with: python tools/generate_excel.py", False),
    ]
    for r, (t, bold) in enumerate(lines, 1):
        c = ws.cell(r, 1, t)
        c.font = Font(name=F, size=13 if bold else 10, bold=bold, color="0A5F5C" if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    # ------------------------------------------------------------------ Data Driven Cases (read by the tests - keep static values only)
    dd = wb.create_sheet("Data Driven Cases")
    header(dd, 1, ["Test Case ID", "Module", "Function Under Test", "Handler", "Technique", "Description", "Test Data File",
                   "Data Columns", "Expected Result Rule", "Owner", "Priority", "Data Rows"])
    for i, r in enumerate(DATA_DRIVEN, 2):
        put(dd, i, list(r) + [csv_rows(r[6])], center=(1, 5, 10, 11, 12))
        dd.cell(i, 1).font = Font(name=F, size=10, bold=True)
    widths(dd, [11, 24, 30, 20, 12, 60, 24, 40, 46, 15, 10, 10])
    dd.freeze_panes = "B2"

    # ------------------------------------------------------------------ Test Inventory
    inv_rows = build_inventory()
    inv = wb.create_sheet("Test Inventory")
    header(inv, 1, ["Test ID", "Module", "Function / Class Under Test", "Test Function", "What it checks", "Technique", "Owner",
                    "Executed Cases", "Markers (groups)", "File"])
    for i, r in enumerate(inv_rows, 2):
        put(inv, i, r, center=(1, 6, 8))
    widths(inv, [10, 26, 40, 58, 70, 12, 15, 10, 40, 52])
    inv.freeze_panes = "E2"
    inv.auto_filter.ref = f"A1:J{len(inv_rows) + 1}"

    # ------------------------------------------------------------------ Coverage Check (formulas)
    cc = wb.create_sheet("Coverage Check")
    cc["A1"] = "Rule: each person writes 5 to 10 automated tests for every function and for every class / module"
    cc["A1"].font = Font(name=F, bold=True, size=11, color="0A5F5C")
    header(cc, 3, ["Module", "Function / Class Under Test", MARYAM, HUDA, "Both", "Status"])
    pairs = []
    for r in inv_rows:
        if (r[1], r[2]) not in pairs:
            pairs.append((r[1], r[2]))
    last = len(inv_rows) + 1
    for i, (mod, fut) in enumerate(pairs, 4):
        put(cc, i, [mod, fut, None, None, None, None], center=(3, 4, 5, 6))
        cc.cell(i, 3, f"=COUNTIFS('Test Inventory'!$B$2:$B${last},$A{i},'Test Inventory'!$C$2:$C${last},$B{i},'Test Inventory'!$G$2:$G${last},C$3)")
        cc.cell(i, 4, f"=COUNTIFS('Test Inventory'!$B$2:$B${last},$A{i},'Test Inventory'!$C$2:$C${last},$B{i},'Test Inventory'!$G$2:$G${last},D$3)")
        cc.cell(i, 5, f"=C{i}+D{i}")
        cc.cell(i, 6, f'=IF(AND(C{i}>=5,C{i}<=10,D{i}>=5,D{i}<=10),"OK","CHECK")')
        for c in range(3, 7):
            cc.cell(i, c).font, cc.cell(i, c).border, cc.cell(i, c).alignment = Font(name=F, size=10), BRD, CEN
    end = 3 + len(pairs)
    t = end + 1
    cc.cell(t, 2, "Total test functions").font = Font(name=F, bold=True, size=10)
    for c in (3, 4, 5):
        col = get_column_letter(c)
        cell = cc.cell(t, c, f"=SUM({col}4:{col}{end})")
        cell.font, cell.border, cell.alignment = Font(name=F, bold=True, size=10), BRD, CEN
    cc.cell(t + 2, 1, "Executed cases (including parametrised variants)").font = Font(name=F, bold=True, size=10)
    for c, who in ((3, MARYAM), (4, HUDA)):
        cc.cell(t + 2, c, f"=SUMIFS('Test Inventory'!$H$2:$H${last},'Test Inventory'!$G$2:$G${last},{get_column_letter(c)}$3)").alignment = CEN
    cc.cell(t + 2, 5, f"=C{t + 2}+D{t + 2}").alignment = CEN
    green = PatternFill("solid", bgColor="DDF3E4", fgColor="DDF3E4")
    red = PatternFill("solid", bgColor="F9DADA", fgColor="F9DADA")
    cc.conditional_formatting.add(f"F4:F{end}", CellIsRule(operator="equal", formula=['"OK"'], fill=green))
    cc.conditional_formatting.add(f"F4:F{end}", CellIsRule(operator="equal", formula=['"CHECK"'], fill=red))
    widths(cc, [30, 46, 16, 16, 10, 10])

    # ------------------------------------------------------------------ Known Defects
    kd = wb.create_sheet("Known Defects")
    header(kd, 1, ["Defect ID", "Description", "Test Function", "Origin"])
    for i, r in enumerate(known_defects(), 2):
        put(kd, i, r, center=(1,))
    widths(kd, [12, 110, 60, 40])

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"written {OUT.relative_to(ROOT)}: {len(DATA_DRIVEN)} data-driven cases, {len(inv_rows)} inventory rows")


if __name__ == "__main__":
    main()
