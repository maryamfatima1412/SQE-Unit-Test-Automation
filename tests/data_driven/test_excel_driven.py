"""Task 3 - data-driven tests.

The test cases are defined in an Excel sheet (docs/Unit_Test_Cases.xlsx, sheet 'Data Driven Cases').
Every case names a CSV file in tests/data. This module reads the Excel sheet, reads the CSV, and turns
every CSV row into one pytest test (parameterization). To add test data, add a row to a CSV file -
no Python code has to change.

If the Excel file is missing, create it with:   python tools/generate_excel.py
"""
import csv
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

import student_portal as sp

ROOT = Path(__file__).resolve().parents[2]
EXCEL = ROOT / "docs" / "Unit_Test_Cases.xlsx"
DATA = ROOT / "tests" / "data"
ME = "2024-001"


# --------------------------------------------------------------------------- #
# Reading the test cases from Excel and the data from CSV
# --------------------------------------------------------------------------- #
def read_cases():
    if not EXCEL.exists():
        raise FileNotFoundError(f"{EXCEL} not found. Run: python tools/generate_excel.py")
    wb = load_workbook(EXCEL, read_only=True, data_only=True)
    rows = list(wb["Data Driven Cases"].iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if r and r[0]]


def read_csv(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def build_params():
    params = []
    for case in read_cases():
        owner_mark = pytest.mark.maryam if str(case["Owner"]).startswith("Maryam") else pytest.mark.huda
        for number, row in enumerate(read_csv(case["Test Data File"]), 1):
            params.append(pytest.param(case["Handler"], row, id=f"{case['Test Case ID']}-row{number:02d}",
                                       marks=[pytest.mark.data_driven, owner_mark]))
    return params


# --------------------------------------------------------------------------- #
# One handler per 'Handler' value in the Excel sheet
# --------------------------------------------------------------------------- #
def handle_grade_from_percentage(row, fx):
    assert sp.grade_from_percentage(float(row["percentage"])) == row["expected_grade"]


def handle_grade_color(row, fx):
    assert sp.grade_color(row["grade"]) == getattr(sp, row["expected_color"])


def handle_overall_attendance(row, fx):
    statuses = [s for s in row["statuses"].split("|") if s]
    records = [fx.att_rec("CS101", f"2026-10-{i:02d}", {ME: s}) for i, s in enumerate(statuses, 1)]
    expected = (int(row["expected_total"]), int(row["expected_present"]), int(row["expected_percent"]))
    assert fx.make_dashboard(roll=ME).overall_attendance(records) == expected


def handle_average_grade(row, fx):
    percentages = [p for p in row["percentages"].split("|") if p]
    approved = [a == "Y" for a in row["approved"].split("|") if a]
    records = [fx.res_rec("CS101", "test", 100, {ME: (0, p, "")}, approved=ok) for p, ok in zip(percentages, approved)]
    grade, count, average = fx.make_dashboard(roll=ME).average_grade(records)
    assert (grade, count) == (row["expected_grade"], int(row["expected_count"]))
    assert average == pytest.approx(float(row["expected_average"]))


def handle_storage_round_trip(row, fx):
    value = json.loads(row["json_value"])
    fx.store.set(row["key"], value)
    assert sp.Storage().get(row["key"]) == value  # a new Storage object = application restart


HANDLERS = {
    "grade_from_percentage": handle_grade_from_percentage,
    "grade_color": handle_grade_color,
    "overall_attendance": handle_overall_attendance,
    "average_grade": handle_average_grade,
    "storage_round_trip": handle_storage_round_trip,
}


@pytest.mark.parametrize("handler, row", build_params())
def test_excel_driven_case(handler, row, make_dashboard, att_rec, res_rec, store):
    """Run one CSV row through the handler named in the Excel sheet."""
    fixtures = SimpleNamespace(make_dashboard=make_dashboard, att_rec=att_rec, res_rec=res_rec, store=store)
    HANDLERS[handler](row, fixtures)


@pytest.mark.data_driven
def test_every_handler_in_excel_exists():
    """Safety net: each 'Handler' value in the Excel sheet must have a function above."""
    assert {c["Handler"] for c in read_cases()} <= set(HANDLERS)


@pytest.mark.data_driven
def test_excel_has_at_least_four_data_driven_cases_and_every_csv_exists():
    """Assignment rule: at least 4 test cases in Excel, each consuming data from an external CSV file."""
    cases = read_cases()
    assert len(cases) >= 4
    assert all((DATA / c["Test Data File"]).exists() for c in cases)
